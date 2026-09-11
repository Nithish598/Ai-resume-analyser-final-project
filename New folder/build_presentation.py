import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_ben10_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Color Palette
    COLOR_BG = RGBColor(6, 11, 17)            # #060B11 Deep Dark Space
    COLOR_CARD = RGBColor(13, 22, 33)         # #0D1621 Glass Slate
    COLOR_CARD_ALT = RGBColor(17, 28, 42)     # #111C2A Elevated Card
    COLOR_GREEN = RGBColor(0, 255, 102)       # #00FF66 Omnitrix Green
    COLOR_CYAN = RGBColor(0, 229, 255)        # #00E5FF Electric Cyan
    COLOR_BLUE = RGBColor(10, 132, 255)       # #0A84FF Hologram Blue
    COLOR_ORANGE = RGBColor(245, 158, 11)     # #F59E0B Warning/Amber
    COLOR_PURPLE = RGBColor(168, 85, 247)     # #A855F7 Violet
    COLOR_WHITE = RGBColor(255, 255, 255)     # Pure White
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# #94A3B8 Slate 400
    COLOR_BORDER_DIM = RGBColor(30, 48, 68)   # #1E3044 Subtle Border
    COLOR_TRACK = RGBColor(22, 36, 52)        # #162434 Progress Track
    
    assets_dir = r"c:\Users\shinc\OneDrive\Documents\AI RECRUITMENT PLATFORM\New folder\assets"
    
    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text, accent_color=COLOR_GREEN):
        tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(4.5), Inches(0.35))
        tag_box.fill.solid()
        tag_box.fill.fore_color.rgb = RGBColor(10, 20, 30)
        tag_box.line.color.rgb = accent_color
        tag_box.line.width = Pt(1)
        tf = tag_box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = tag_text
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = accent_color
        p.font.name = "Consolas"
        p.alignment = PP_ALIGN.CENTER
        
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.95))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE
        p_title.font.name = "Segoe UI"
        
        p_sub = tf_title.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.font.name = "Segoe UI"
        
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = accent_color
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_BORDER_DIM, border_width=1.2):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
        return card

    def add_stat_gauge(slide, left, top, width, height, score_text, label_text, badge_text, color=COLOR_GREEN):
        card = add_card(slide, left, top, width, height, COLOR_CARD, color, 1.8)
        
        c_size = 2.4
        c_left = left + (width - c_size) / 2
        c_top = top + 0.35
        
        outer_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_left), Inches(c_top), Inches(c_size), Inches(c_size))
        outer_circle.fill.solid()
        outer_circle.fill.fore_color.rgb = RGBColor(12, 24, 36)
        outer_circle.line.color.rgb = color
        outer_circle.line.width = Pt(3)
        
        inner_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_left + 0.15), Inches(c_top + 0.15), Inches(c_size - 0.3), Inches(c_size - 0.3))
        inner_circle.fill.solid()
        inner_circle.fill.fore_color.rgb = RGBColor(6, 12, 18)
        inner_circle.line.color.rgb = COLOR_BORDER_DIM
        inner_circle.line.width = Pt(1)
        
        tb = slide.shapes.add_textbox(Inches(c_left), Inches(c_top + 0.5), Inches(c_size), Inches(1.3))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = score_text
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.alignment = PP_ALIGN.CENTER
        
        tb_label = slide.shapes.add_textbox(Inches(left + 0.2), Inches(c_top + c_size + 0.15), Inches(width - 0.4), Inches(0.6))
        tf_l = tb_label.text_frame
        p_l = tf_l.paragraphs[0]
        p_l.text = label_text
        p_l.font.size = Pt(13)
        p_l.font.bold = True
        p_l.font.color.rgb = COLOR_WHITE
        p_l.font.name = "Segoe UI"
        p_l.alignment = PP_ALIGN.CENTER
        
        if badge_text:
            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + (width - 2.5) / 2), Inches(c_top + c_size + 0.8), Inches(2.5), Inches(0.38))
            badge.fill.solid()
            badge.fill.fore_color.rgb = RGBColor(0, 40, 20)
            badge.line.color.rgb = color
            badge.line.width = Pt(1.5)
            b_tf = badge.text_frame
            b_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            bp = b_tf.paragraphs[0]
            bp.text = badge_text
            bp.font.size = Pt(11)
            bp.font.bold = True
            bp.font.color.rgb = color
            bp.font.name = "Segoe UI"
            bp.alignment = PP_ALIGN.CENTER

    def add_progress_row(slide, left, top, width, label, percent, bar_color=COLOR_GREEN):
        lbl_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width - 1.0), Inches(0.28))
        lbl_tf = lbl_box.text_frame
        lbl_tf.margin_left = lbl_tf.margin_top = lbl_tf.margin_right = lbl_tf.margin_bottom = 0
        lp = lbl_tf.paragraphs[0]
        lp.text = label
        lp.font.size = Pt(11.5)
        lp.font.color.rgb = COLOR_WHITE
        lp.font.bold = True
        lp.font.name = "Segoe UI"
        
        pct_box = slide.shapes.add_textbox(Inches(left + width - 1.0), Inches(top), Inches(1.0), Inches(0.28))
        pct_tf = pct_box.text_frame
        pct_tf.margin_left = pct_tf.margin_top = pct_tf.margin_right = pct_tf.margin_bottom = 0
        pp = pct_tf.paragraphs[0]
        pp.text = f"{percent}%"
        pp.font.size = Pt(12)
        pp.font.bold = True
        pp.font.color.rgb = bar_color
        pp.font.name = "Consolas"
        pp.alignment = PP_ALIGN.RIGHT
        
        track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top + 0.32), Inches(width), Inches(0.14))
        track.fill.solid()
        track.fill.fore_color.rgb = COLOR_TRACK
        track.line.fill.background()
        
        fill_w = max(0.1, width * (percent / 100.0))
        fill_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top + 0.32), Inches(fill_w), Inches(0.14))
        fill_bar.fill.solid()
        fill_bar.fill.fore_color.rgb = bar_color
        fill_bar.line.fill.background()

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide1)
    
    ben_img_path = os.path.join(assets_dir, "ben_hero.png")
    if os.path.exists(ben_img_path):
        # 896x1200 aspect ratio (0.7467). Setting height to 6.1 inches gives width 4.55 inches, top 0.7 gives 0.7 margin top & bottom
        slide1.shapes.add_picture(ben_img_path, Inches(0.7), Inches(0.7), Inches(4.55), Inches(6.1))
    
    add_card(slide1, 5.7, 0.7, 7.0, 6.1, COLOR_CARD, COLOR_GREEN, 1.8)
    
    top_badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.0), Inches(1.0), Inches(6.4), Inches(0.38))
    top_badge.fill.solid()
    top_badge.fill.fore_color.rgb = RGBColor(0, 30, 20)
    top_badge.line.color.rgb = COLOR_GREEN
    top_badge.line.width = Pt(1)
    tb_badge = top_badge.text_frame
    tb_badge.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp1 = tb_badge.paragraphs[0]
    bp1.text = "⚡ OMNITRIX NEURAL INTERFACE // ACTIVE v2.4 // GALVAN AI"
    bp1.font.size = Pt(9.5)
    bp1.font.bold = True
    bp1.font.color.rgb = COLOR_GREEN
    bp1.font.name = "Consolas"
    bp1.alignment = PP_ALIGN.CENTER
    
    tb_main = slide1.shapes.add_textbox(Inches(6.0), Inches(1.6), Inches(6.4), Inches(1.6))
    tf_main = tb_main.text_frame
    tf_main.word_wrap = True
    p1 = tf_main.paragraphs[0]
    p1.text = "AI RECRUITMENT"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.font.name = "Segoe UI"
    
    p2 = tf_main.add_paragraph()
    p2.text = "PLATFORM"
    p2.font.size = Pt(40)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_GREEN
    p2.font.name = "Segoe UI"
    
    sub_pill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.0), Inches(3.4), Inches(3.2), Inches(0.42))
    sub_pill.fill.solid()
    sub_pill.fill.fore_color.rgb = RGBColor(0, 45, 60)
    sub_pill.line.color.rgb = COLOR_CYAN
    sub_pill.line.width = Pt(1.5)
    sp_tf = sub_pill.text_frame
    sp_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    spp = sp_tf.paragraphs[0]
    spp.text = "ADVANCED MODULES"
    spp.font.size = Pt(11)
    spp.font.bold = True
    spp.font.color.rgb = COLOR_CYAN
    spp.font.name = "Consolas"
    spp.alignment = PP_ALIGN.CENTER
    
    tb_tag = slide1.shapes.add_textbox(Inches(6.0), Inches(3.95), Inches(6.4), Inches(0.5))
    tf_tag = tb_tag.text_frame
    ptag = tf_tag.paragraphs[0]
    ptag.text = "Smarter Hiring  •  Better Matches  •  Brighter Futures"
    ptag.font.size = Pt(13)
    ptag.font.bold = True
    ptag.font.color.rgb = COLOR_WHITE
    ptag.font.name = "Segoe UI"
    
    chips = [
        ("🎯 78% Interview Readiness", COLOR_CYAN),
        ("⚡ 84% Candidate Role Fit", COLOR_GREEN),
        ("💰 Dynamic Salary Prediction", COLOR_ORANGE),
        ("🌐 Neural Ingestion Pipeline", COLOR_PURPLE),
    ]
    for i, (chip_txt, chip_col) in enumerate(chips):
        cx = 6.0 + (i % 2) * 3.25
        cy = 4.55 + (i // 2) * 0.7
        c_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(cy), Inches(3.1), Inches(0.55))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = COLOR_CARD_ALT
        c_shape.line.color.rgb = chip_col
        c_shape.line.width = Pt(1)
        c_tf = c_shape.text_frame
        c_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp = c_tf.paragraphs[0]
        cp.text = chip_txt
        cp.font.size = Pt(10)
        cp.font.bold = True
        cp.font.color.rgb = COLOR_WHITE
        cp.font.name = "Segoe UI"
        cp.alignment = PP_ALIGN.CENTER

    foot = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.0), Inches(6.05), Inches(6.4), Inches(0.48))
    foot.fill.solid()
    foot.fill.fore_color.rgb = RGBColor(10, 16, 24)
    foot.line.color.rgb = COLOR_BORDER_DIM
    foot.line.width = Pt(1)
    f_tf = foot.text_frame
    f_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fp = f_tf.paragraphs[0]
    fp.text = "Team D  |  AI/ML Internship  |  RynixSoft"
    fp.font.size = Pt(11)
    fp.font.bold = True
    fp.font.color.rgb = COLOR_TEXT_MUTED
    fp.font.name = "Segoe UI"
    fp.alignment = PP_ALIGN.CENTER

    notes1 = slide1.notes_slide.notes_text_frame
    notes1.text = (
        "[CINEMATIC OPENING ANIMATION]\n"
        "- Camera starts in total darkness, slowly panning in on 3D Ben Tennyson.\n"
        "- Ben raises his arm; the Omnitrix dial charges with green particle lasers and rotating holographic HUD rings.\n"
        "- Ben slams the dial: A burst of emerald light projects the giant 'AI RECRUITMENT PLATFORM' holographic HUD in mid-air.\n"
        "- SOUND: Low hum of alien energy -> high-pitch capacitor charge -> futuristic Omnitrix activation shockwave."
    )

    # =========================================================================
    # SLIDE 2: PROJECT OVERVIEW (4-Stage Pipeline)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2)
    add_header(slide2, "⚡ MODULE 01 // SYSTEM ARCHITECTURE", "PROJECT OVERVIEW — 4-STAGE AI PIPELINE", 
               "AI-powered platform to make modern talent recruitment smarter, faster, and demonstrably fairer.")
    
    stages = [
        ("01", "RESUME", "Upload & Ingestion", "• Multi-format PDF/DOCX\n• Text tokenization\n• Layout preservation", COLOR_CYAN),
        ("02", "CANDIDATE PROFILE", "Extract & Structure", "• NLP Entity Extraction\n• Standardized Taxonomy\n• Experience indexing", COLOR_BLUE),
        ("03", "AI PREDICTIONS", "Analyze & Score", "• Interview readiness\n• Role fit probability\n• Salary market band", COLOR_GREEN),
        ("04", "HIRING DECISION", "Match & Placement", "• Objective ranking\n• Gap-free evaluation\n• High retention placement", COLOR_ORANGE),
    ]
    
    card_w = 2.6
    card_h = 2.45
    card_y = 2.1
    for i, (num, title, sub, desc, col) in enumerate(stages):
        cx = 0.8 + i * 3.05
        c = add_card(slide2, cx, card_y, card_w, card_h, COLOR_CARD, col, 1.5)
        
        nb = slide2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.2), Inches(card_y + 0.2), Inches(0.45), Inches(0.45))
        nb.fill.solid()
        nb.fill.fore_color.rgb = col
        nb.line.fill.background()
        n_tf = nb.text_frame
        n_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        np = n_tf.paragraphs[0]
        np.text = num
        np.font.size = Pt(11)
        np.font.bold = True
        np.font.color.rgb = RGBColor(0, 0, 0)
        np.font.name = "Consolas"
        np.alignment = PP_ALIGN.CENTER
        
        tb = slide2.shapes.add_textbox(Inches(cx + 0.75), Inches(card_y + 0.15), Inches(card_w - 0.9), Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.name = "Segoe UI"
        
        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = col
        p_sub.font.name = "Segoe UI"
        
        tb_desc = slide2.shapes.add_textbox(Inches(cx + 0.2), Inches(card_y + 0.85), Inches(card_w - 0.4), Inches(1.4))
        tf_d = tb_desc.text_frame
        tf_d.word_wrap = True
        tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
        dp = tf_d.paragraphs[0]
        dp.text = desc
        dp.font.size = Pt(10)
        dp.font.color.rgb = COLOR_TEXT_MUTED
        dp.font.name = "Segoe UI"
        
        if i < 3:
            arrow = slide2.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + card_w + 0.08), Inches(card_y + card_h / 2 - 0.15), Inches(0.28), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COLOR_GREEN
            arrow.line.fill.background()

    add_card(slide2, 0.8, 4.8, 7.8, 2.3, COLOR_CARD, COLOR_BORDER_DIM, 1.2)
    
    h_tb = slide2.shapes.add_textbox(Inches(1.1), Inches(4.95), Inches(7.2), Inches(0.4))
    h_tf = h_tb.text_frame
    hp = h_tf.paragraphs[0]
    hp.text = "HOW AI & MACHINE LEARNING REVOLUTIONIZES HIRING"
    hp.font.size = Pt(13)
    hp.font.bold = True
    hp.font.color.rgb = COLOR_GREEN
    hp.font.name = "Segoe UI"
    
    benefits = [
        ("✓  Finds the Best Matches:", "Deep semantic embedding goes far beyond basic keyword searches to match true candidate capability."),
        ("✓  Predicts Performance & Success:", "Empirical multi-variate machine learning models forecast interview readiness & organizational fit."),
        ("✓  Provides Data-Driven Transparency:", "Replaces subjective gut-feel with verifiable metrics, structured skill gaps, and market salary bands."),
    ]
    for b_idx, (b_title, b_desc) in enumerate(benefits):
        b_box = slide2.shapes.add_textbox(Inches(1.1), Inches(5.35 + b_idx * 0.55), Inches(7.2), Inches(0.5))
        b_tf = b_box.text_frame
        b_tf.word_wrap = True
        b_tf.margin_left = b_tf.margin_top = b_tf.margin_right = b_tf.margin_bottom = 0
        bp_t = b_tf.paragraphs[0]
        bp_t.text = f"{b_title} "
        bp_t.font.size = Pt(10.5)
        bp_t.font.bold = True
        bp_t.font.color.rgb = COLOR_WHITE
        bp_t.font.name = "Segoe UI"
        
        run = bp_t.add_run()
        run.text = b_desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED
        run.font.size = Pt(10.5)

    core_img_path = os.path.join(assets_dir, "omnitrix_core.png")
    if os.path.exists(core_img_path):
        slide2.shapes.add_picture(core_img_path, Inches(8.9), Inches(4.75), Inches(3.6), Inches(2.4))
        
    notes2 = slide2.notes_slide.notes_text_frame
    notes2.text = (
        "[TRANSITION CUE: SLIDE 1 -> SLIDE 2]\n"
        "- Camera zooms directly through the center of Ben's glowing Omnitrix dial.\n"
        "- 4 glowing 3D glass cards materialize in sequence from left to right along an electric green particle conveyor.\n"
        "- Omnitrix Core HUD on the bottom-right projects scanning laser lines across each stage in real time."
    )

    # =========================================================================
    # SLIDE 3: INTERVIEW PERFORMANCE PREDICTION
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide3)
    add_header(slide3, "⚡ MODULE 02 // PREDICTIVE READINESS ENGINE", "INTERVIEW PERFORMANCE PREDICTION",
               "Multi-factor predictive model estimating candidate readiness, technical depth, and interview success probability.")
    
    add_card(slide3, 0.8, 2.1, 3.5, 4.9, COLOR_CARD, COLOR_BORDER_DIM, 1.2)
    in_tb = slide3.shapes.add_textbox(Inches(1.05), Inches(2.25), Inches(3.0), Inches(0.4))
    in_tf = in_tb.text_frame
    inp = in_tf.paragraphs[0]
    inp.text = "KEY INPUT SIGNALS"
    inp.font.size = Pt(13)
    inp.font.bold = True
    inp.font.color.rgb = COLOR_CYAN
    inp.font.name = "Segoe UI"
    
    inputs_s3 = [
        ("🎯", "Skills & Technical Knowledge", "Deep assessment of syntax, frameworks & problem domain"),
        ("💼", "Experience & Education Tier", "Verified years of practice, accredited pedigree"),
        ("🚀", "Projects & Certifications", "Hands-on repo quality, cloud & security certs"),
        ("📈", "Previous Performance History", "Historical screening milestones and peer ratings"),
        ("🗣️", "Interview-Related Test Scores", "Aptitude, quantitative logic & system intuition"),
    ]
    for idx, (icon, title, sub) in enumerate(inputs_s3):
        box = slide3.shapes.add_textbox(Inches(1.05), Inches(2.7 + idx * 0.82), Inches(3.0), Inches(0.75))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"{icon}  {title}"
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.name = "Segoe UI"
        
        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(9)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.font.name = "Segoe UI"

    add_stat_gauge(slide3, 4.6, 2.1, 4.5, 2.7, "78%", "ESTIMATED INTERVIEW READINESS", "★ GREAT POTENTIAL! ★", COLOR_GREEN)
    
    add_card(slide3, 4.6, 5.0, 4.5, 2.0, COLOR_CARD, COLOR_BORDER_DIM, 1.2)
    b_head = slide3.shapes.add_textbox(Inches(4.8), Inches(5.1), Inches(4.1), Inches(0.3))
    b_htf = b_head.text_frame
    bhp = b_htf.paragraphs[0]
    bhp.text = "READINESS BREAKDOWN METRICS"
    bhp.font.size = Pt(11)
    bhp.font.bold = True
    bhp.font.color.rgb = COLOR_GREEN
    bhp.font.name = "Segoe UI"
    
    add_progress_row(slide3, 4.8, 5.45, 4.1, "Technical Knowledge", 85, COLOR_CYAN)
    add_progress_row(slide3, 4.8, 6.0, 4.1, "Communication Clarity", 75, COLOR_GREEN)
    add_progress_row(slide3, 4.8, 6.55, 4.1, "Problem Solving Aptitude", 70, COLOR_BLUE)

    gm_img_path = os.path.join(assets_dir, "grey_matter.png")
    if os.path.exists(gm_img_path):
        slide3.shapes.add_picture(gm_img_path, Inches(9.35), Inches(2.1), Inches(3.2), Inches(3.2))
        
    add_card(slide3, 9.35, 5.4, 3.2, 1.6, COLOR_CARD_ALT, COLOR_CYAN, 1.2)
    rec_box = slide3.shapes.add_textbox(Inches(9.5), Inches(5.5), Inches(2.9), Inches(1.4))
    r_tf = rec_box.text_frame
    r_tf.word_wrap = True
    r_tf.margin_left = r_tf.margin_top = r_tf.margin_right = r_tf.margin_bottom = 0
    rp = r_tf.paragraphs[0]
    rp.text = "FOCUS AREAS FOR GROWTH"
    rp.font.size = Pt(10.5)
    rp.font.bold = True
    rp.font.color.rgb = COLOR_CYAN
    rp.font.name = "Segoe UI"
    
    recs = [
        "• Advanced distributed system design",
        "• High-pressure live coding speed",
        "• Architectural trade-off articulation"
    ]
    for r in recs:
        p_r = r_tf.add_paragraph()
        p_r.text = r
        p_r.font.size = Pt(9.5)
        p_r.font.color.rgb = COLOR_TEXT_MUTED
        p_r.font.name = "Segoe UI"

    notes3 = slide3.notes_slide.notes_text_frame
    notes3.text = (
        "[TRANSITION CUE: SLIDE 2 -> SLIDE 3]\n"
        "- Omnitrix fires an emerald scanning fan beam.\n"
        "- The 3D candidate wireframe materializes into the glowing 78% Readiness radial meter.\n"
        "- 3D Grey Matter appears on the right, dynamically manipulating neural nodes and readiness breakdown charts."
    )

    # =========================================================================
    # SLIDE 4: CANDIDATE SUCCESS PREDICTION
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide4)
    add_header(slide4, "⚡ MODULE 03 // ROLE FIT & SUCCESS ENGINE", "CANDIDATE SUCCESS PREDICTION",
               "Estimates how well a candidate aligns with the role requirements and predicts long-term retention and success.")
    
    add_stat_gauge(slide4, 0.8, 2.1, 4.4, 2.7, "84%", "ESTIMATED ROLE FIT", "★ GREAT MATCH! (TOP 5%) ★", COLOR_CYAN)
    
    add_card(slide4, 0.8, 5.0, 4.4, 2.0, COLOR_CARD, COLOR_BORDER_DIM, 1.2)
    af_box = slide4.shapes.add_textbox(Inches(1.0), Inches(5.1), Inches(4.0), Inches(0.3))
    af_tf = af_box.text_frame
    afp = af_tf.paragraphs[0]
    afp.text = "CORE ALIGNMENT BREAKDOWN"
    afp.font.size = Pt(11)
    afp.font.bold = True
    afp.font.color.rgb = COLOR_CYAN
    afp.font.name = "Segoe UI"
    
    add_progress_row(slide4, 1.0, 5.42, 4.0, "Skill Alignment", 90, COLOR_GREEN)
    add_progress_row(slide4, 1.0, 5.92, 4.0, "Experience Alignment", 85, COLOR_CYAN)
    add_progress_row(slide4, 1.0, 6.42, 4.0, "JD Semantic Match", 82, COLOR_BLUE)

    add_card(slide4, 5.45, 2.1, 4.2, 2.35, COLOR_CARD, COLOR_GREEN, 1.5)
    s_tb = slide4.shapes.add_textbox(Inches(5.65), Inches(2.25), Inches(3.8), Inches(0.4))
    s_tf = s_tb.text_frame
    sp = s_tf.paragraphs[0]
    sp.text = "VERIFIED CORE STRENGTHS"
    sp.font.size = Pt(12)
    sp.font.bold = True
    sp.font.color.rgb = COLOR_GREEN
    sp.font.name = "Segoe UI"
    
    strengths = [
        ("✓", "Exceptional Python & Async stack proficiency"),
        ("✓", "Solid relational SQL schema & query optimization"),
        ("✓", "Proven hands-on backend microservices architecture"),
        ("✓", "Strong algorithmic baseline and clean modular code")
    ]
    for s_icon, s_txt in strengths:
        p_s = s_tf.add_paragraph()
        p_s.text = f"{s_icon}  {s_txt}"
        p_s.font.size = Pt(9.5)
        p_s.font.color.rgb = COLOR_WHITE
        p_s.font.name = "Segoe UI"
        
    add_card(slide4, 5.45, 4.65, 4.2, 2.35, COLOR_CARD, COLOR_ORANGE, 1.5)
    i_tb = slide4.shapes.add_textbox(Inches(5.65), Inches(4.8), Inches(3.8), Inches(0.4))
    i_tf = i_tb.text_frame
    ip = i_tf.paragraphs[0]
    ip.text = "TARGETED IMPROVEMENT AREAS"
    ip.font.size = Pt(12)
    ip.font.bold = True
    ip.font.color.rgb = COLOR_ORANGE
    ip.font.name = "Segoe UI"
    
    improvements = [
        ("⚠", "Limited enterprise AWS/GCP cloud orchestration"),
        ("⚠", "Minimal production containerization (Kubernetes)"),
        ("⚠", "CI/CD automated regression pipeline configuration"),
        ("⚠", "Needs mentoring on distributed log monitoring")
    ]
    for i_icon, i_txt in improvements:
        p_i = i_tf.add_paragraph()
        p_i.text = f"{i_icon}  {i_txt}"
        p_i.font.size = Pt(9.5)
        p_i.font.color.rgb = COLOR_TEXT_MUTED
        p_i.font.name = "Segoe UI"

    xlr8_img_path = os.path.join(assets_dir, "xlr8.png")
    if os.path.exists(xlr8_img_path):
        slide4.shapes.add_picture(xlr8_img_path, Inches(9.85), Inches(2.1), Inches(2.8), Inches(2.8))
        
    add_card(slide4, 9.85, 5.0, 2.8, 2.0, COLOR_CARD_ALT, COLOR_CYAN, 1.2)
    spd_box = slide4.shapes.add_textbox(Inches(10.0), Inches(5.15), Inches(2.5), Inches(1.7))
    spd_tf = spd_box.text_frame
    spd_tf.word_wrap = True
    spd_tf.margin_left = spd_tf.margin_top = spd_tf.margin_right = spd_tf.margin_bottom = 0
    spdp = spd_tf.paragraphs[0]
    spdp.text = "XLR8 RAPID MATCHING"
    spdp.font.size = Pt(10.5)
    spdp.font.bold = True
    spdp.font.color.rgb = COLOR_CYAN
    spdp.font.name = "Segoe UI"
    
    spd_stats = [
        "⚡ Match Time: < 0.4s",
        "⚡ Cohort Size: 10,000+",
        "⚡ Retention Rate: 93.4%",
        "⚡ Bias Divergence: 0.02%"
    ]
    for s in spd_stats:
        p_spd = spd_tf.add_paragraph()
        p_spd.text = s
        p_spd.font.size = Pt(9.5)
        p_spd.font.color.rgb = COLOR_TEXT_MUTED
        p_spd.font.name = "Consolas"

    notes4 = slide4.notes_slide.notes_text_frame
    notes4.text = (
        "[TRANSITION CUE: SLIDE 3 -> SLIDE 4]\n"
        "- The 78% Readiness dial morphs smoothly into an 84% Role Fit holographic ring.\n"
        "- XLR8 dashes into the frame leaving electric cyan speed trails.\n"
        "- The Strengths and Improvement glass panels slide outward from XLR8's velocity trails."
    )

    # =========================================================================
    # SLIDE 5: SALARY PREDICTION
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide5)
    add_header(slide5, "⚡ MODULE 04 // COMPENSATION MODELING ENGINE", "SALARY PREDICTION — FAIR & COMPETITIVE",
               "Multi-variable regression model calculating objective, market-aligned compensation benchmarks.")
    
    add_card(slide5, 0.8, 2.1, 3.6, 4.9, COLOR_CARD, COLOR_BORDER_DIM, 1.2)
    s_in_tb = slide5.shapes.add_textbox(Inches(1.05), Inches(2.25), Inches(3.1), Inches(0.4))
    s_in_tf = s_in_tb.text_frame
    s_inp = s_in_tf.paragraphs[0]
    s_inp.text = "COMPENSATION PARAMETERS"
    s_inp.font.size = Pt(12.5)
    s_inp.font.bold = True
    s_inp.font.color.rgb = COLOR_ORANGE
    s_inp.font.name = "Segoe UI"
    
    salary_inputs = [
        ("🏢", "Job Role & Seniority Tier", "Junior, Mid, Senior, Staff, Principal"),
        ("📈", "Experience & Tech Stack", "Years of tenure & scarcity of core skills"),
        ("📍", "Geographic Location", "Tier-1 vs Tier-2 cost-of-living adjustments"),
        ("🎓", "Education & Academic Pedigree", "B.Tech/M.S. & specialized research"),
        ("🏭", "Industry Domain & Sector", "FinTech, HealthTech, AI, Enterprise SaaS"),
        ("📜", "Verified Certifications", "AWS Solutions Architect, CKA, PMP"),
        ("💼", "Employment Model", "Full-time, Contract, Remote, Hybrid"),
    ]
    for s_idx, (s_ic, s_ti, s_su) in enumerate(salary_inputs):
        s_box = slide5.shapes.add_textbox(Inches(1.05), Inches(2.7 + s_idx * 0.58), Inches(3.1), Inches(0.55))
        stf = s_box.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
        p = stf.paragraphs[0]
        p.text = f"{s_ic}  {s_ti}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.name = "Segoe UI"
        
        p_sub = stf.add_paragraph()
        p_sub.text = s_su
        p_sub.font.size = Pt(8.5)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.font.name = "Segoe UI"

    add_card(slide5, 4.7, 2.1, 4.7, 4.9, COLOR_CARD, COLOR_ORANGE, 1.8)
    
    sal_head = slide5.shapes.add_textbox(Inches(4.9), Inches(2.25), Inches(4.3), Inches(0.4))
    sh_tf = sal_head.text_frame
    shp = sh_tf.paragraphs[0]
    shp.text = "ESTIMATED MARKET SALARY RANGE"
    shp.font.size = Pt(12.5)
    shp.font.bold = True
    shp.font.color.rgb = COLOR_ORANGE
    shp.font.name = "Segoe UI"
    shp.alignment = PP_ALIGN.CENTER
    
    val_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.0), Inches(2.75), Inches(4.1), Inches(1.1))
    val_box.fill.solid()
    val_box.fill.fore_color.rgb = RGBColor(20, 15, 8)
    val_box.line.color.rgb = COLOR_ORANGE
    val_box.line.width = Pt(2)
    v_tf = val_box.text_frame
    v_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    vp = v_tf.paragraphs[0]
    vp.text = "₹6 LPA — ₹9 LPA"
    vp.font.size = Pt(30)
    vp.font.bold = True
    vp.font.color.rgb = COLOR_ORANGE
    vp.font.name = "Segoe UI"
    vp.alignment = PP_ALIGN.CENTER
    
    tiers = [
        ("LOWER BOUND", "₹6.0 LPA", COLOR_CYAN),
        ("EXPECTED TARGET", "₹7.5 LPA", COLOR_GREEN),
        ("UPPER CEILING", "₹9.0 LPA", COLOR_ORANGE),
    ]
    for t_i, (t_label, t_val, t_c) in enumerate(tiers):
        t_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.0 + t_i * 1.4), Inches(4.0), Inches(1.3), Inches(0.85))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = COLOR_CARD_ALT
        t_card.line.color.rgb = t_c
        t_card.line.width = Pt(1)
        tt_tf = t_card.text_frame
        ttp = tt_tf.paragraphs[0]
        ttp.text = t_label
        ttp.font.size = Pt(7.5)
        ttp.font.bold = True
        ttp.font.color.rgb = COLOR_TEXT_MUTED
        ttp.font.name = "Consolas"
        ttp.alignment = PP_ALIGN.CENTER
        
        ttp2 = tt_tf.add_paragraph()
        ttp2.text = t_val
        ttp2.font.size = Pt(12)
        ttp2.font.bold = True
        ttp2.font.color.rgb = t_c
        ttp2.font.name = "Segoe UI"
        ttp2.alignment = PP_ALIGN.CENTER

    dist_tb = slide5.shapes.add_textbox(Inches(5.0), Inches(5.0), Inches(4.1), Inches(0.3))
    dt_tf = dist_tb.text_frame
    dp_h = dt_tf.paragraphs[0]
    dp_h.text = "MARKET DISTRIBUTION PERCENTILES"
    dp_h.font.size = Pt(10)
    dp_h.font.bold = True
    dp_h.font.color.rgb = COLOR_WHITE
    dp_h.font.name = "Segoe UI"
    
    bars = [
        ("25th % (Entry)", 45, COLOR_CYAN),
        ("50th % (Median)", 70, COLOR_GREEN),
        ("75th % (Target)", 85, COLOR_ORANGE),
        ("90th % (Top Tier)", 95, COLOR_PURPLE),
    ]
    for b_i, (b_name, b_pct, b_col) in enumerate(bars):
        add_progress_row(slide5, 5.0, 5.35 + b_i * 0.38, 4.1, b_name, b_pct, b_col)

    dh_img_path = os.path.join(assets_dir, "diamondhead.png")
    if os.path.exists(dh_img_path):
        slide5.shapes.add_picture(dh_img_path, Inches(9.7), Inches(2.1), Inches(2.9), Inches(2.9))
        
    add_card(slide5, 9.7, 5.15, 2.9, 1.85, COLOR_CARD_ALT, COLOR_GREEN, 1.2)
    vt_box = slide5.shapes.add_textbox(Inches(9.85), Inches(5.3), Inches(2.6), Inches(1.5))
    vt_tf = vt_box.text_frame
    vt_tf.word_wrap = True
    vt_tf.margin_left = vt_tf.margin_top = vt_tf.margin_right = vt_tf.margin_bottom = 0
    vtp = vt_tf.paragraphs[0]
    vtp.text = "FAIR & UNBIASED PRICING"
    vtp.font.size = Pt(10.5)
    vtp.font.bold = True
    vtp.font.color.rgb = COLOR_GREEN
    vtp.font.name = "Segoe UI"
    
    vt_points = [
        "✓ 50,000+ benchmark salaries",
        "✓ Zero gender/pedigree bias",
        "✓ Real-time tech stack inflation",
        "✓ Win-win retention guarantee"
    ]
    for p in vt_points:
        p_vt = vt_tf.add_paragraph()
        p_vt.text = p
        p_vt.font.size = Pt(9.5)
        p_vt.font.color.rgb = COLOR_TEXT_MUTED
        p_vt.font.name = "Segoe UI"

    notes5 = slide5.notes_slide.notes_text_frame
    notes5.text = (
        "[TRANSITION CUE: SLIDE 4 -> SLIDE 5]\n"
        "- Omnitrix rotates dial with metallic click sounds.\n"
        "- Gold and emerald coins/data particles burst outward, coalescing into the ₹6 LPA — ₹9 LPA hero card.\n"
        "- Diamondhead creates solid crystalline data pillars representing robust, unbiased compensation standards."
    )

    # =========================================================================
    # SLIDE 6: HOW THE ADVANCED MODULES WORK TOGETHER
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide6)
    add_header(slide6, "⚡ MODULE 05 // SYSTEM INTEGRATION & PIPELINE", "HOW THE ADVANCED MODULES WORK TOGETHER",
               "The complete end-to-end recruitment intelligence pipeline: from data ingestion to predictive insight engines.")
    
    add_card(slide6, 0.8, 2.1, 7.8, 4.9, COLOR_CARD, COLOR_GREEN, 1.8)
    
    flow_head = slide6.shapes.add_textbox(Inches(1.1), Inches(2.25), Inches(7.2), Inches(0.35))
    fh_tf = flow_head.text_frame
    fhp = fh_tf.paragraphs[0]
    fhp.text = "CORE GALVAN NEURAL PIPELINE — ALL-IN-ONE RECRUITMENT FLOW"
    fhp.font.size = Pt(13)
    fhp.font.bold = True
    fhp.font.color.rgb = COLOR_GREEN
    fhp.font.name = "Segoe UI"
    
    flow_steps = [
        ("01", "Resume Processing", "Extracts, normalizes, and tokenizes candidate experience & skills", COLOR_CYAN),
        ("02", "Job Role Recommendation", "Maps multi-dimensional profiles to optimal organizational roles", COLOR_BLUE),
        ("03", "Skill Gap Analysis", "Pinpoints missing competencies and calculates development effort", COLOR_PURPLE),
        ("04", "JD Matching & Candidate Ranking", "Performs semantic scoring against live open job requisitions", COLOR_GREEN),
    ]
    for f_idx, (f_no, f_ti, f_de, f_col) in enumerate(flow_steps):
        s_top = 2.75 + f_idx * 0.95
        sc = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(s_top), Inches(3.6), Inches(0.8))
        sc.fill.solid()
        sc.fill.fore_color.rgb = COLOR_CARD_ALT
        sc.line.color.rgb = f_col
        sc.line.width = Pt(1.2)
        
        nb = slide6.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(s_top + 0.15), Inches(0.5), Inches(0.5))
        nb.fill.solid()
        nb.fill.fore_color.rgb = f_col
        nb.line.fill.background()
        nb_tf = nb.text_frame
        nb_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        nbp = nb_tf.paragraphs[0]
        nbp.text = f_no
        nbp.font.size = Pt(11)
        nbp.font.bold = True
        nbp.font.color.rgb = RGBColor(0, 0, 0)
        nbp.font.name = "Consolas"
        nbp.alignment = PP_ALIGN.CENTER
        
        stb = slide6.shapes.add_textbox(Inches(1.8), Inches(s_top + 0.1), Inches(2.8), Inches(0.65))
        stf = stb.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
        stp = stf.paragraphs[0]
        stp.text = f_ti
        stp.font.size = Pt(11)
        stp.font.bold = True
        stp.font.color.rgb = COLOR_WHITE
        stp.font.name = "Segoe UI"
        
        st_sub = stf.add_paragraph()
        st_sub.text = f_de
        st_sub.font.size = Pt(8.5)
        st_sub.font.color.rgb = COLOR_TEXT_MUTED
        st_sub.font.name = "Segoe UI"
        
        if f_idx < 3:
            arr = slide6.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(1.35), Inches(s_top + 0.8), Inches(0.2), Inches(0.15))
            arr.fill.solid()
            arr.fill.fore_color.rgb = COLOR_GREEN
            arr.line.fill.background()

    fan_tb = slide6.shapes.add_textbox(Inches(4.9), Inches(2.75), Inches(3.5), Inches(0.35))
    ft_tf = fan_tb.text_frame
    ftp = ft_tf.paragraphs[0]
    ftp.text = "INTEGRATED PREDICTION ENGINES"
    ftp.font.size = Pt(11)
    ftp.font.bold = True
    ftp.font.color.rgb = COLOR_CYAN
    ftp.font.name = "Segoe UI"
    
    outputs = [
        ("INTERVIEW READINESS PREDICTOR", "78% Probability Score", "Estimates candidate interview performance & technical competency breakdown", COLOR_CYAN),
        ("CANDIDATE SUCCESS PREDICTOR", "84% Long-Term Role Fit", "Forecasts long-term organization alignment, cultural match, and retention", COLOR_GREEN),
        ("SALARY VALUATION ENGINE", "₹6 LPA — ₹9 LPA Benchmark", "Calculates competitive, fair, and multi-variable market compensation", COLOR_ORANGE),
    ]
    for o_idx, (o_ti, o_sub, o_de, o_col) in enumerate(outputs):
        o_top = 3.2 + o_idx * 1.15
        oc = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.9), Inches(o_top), Inches(3.5), Inches(1.0))
        oc.fill.solid()
        oc.fill.fore_color.rgb = COLOR_CARD_ALT
        oc.line.color.rgb = o_col
        oc.line.width = Pt(1.5)
        
        otb = slide6.shapes.add_textbox(Inches(5.05), Inches(o_top + 0.1), Inches(3.2), Inches(0.8))
        otf = otb.text_frame
        otf.word_wrap = True
        otf.margin_left = otf.margin_top = otf.margin_right = otf.margin_bottom = 0
        otp = otf.paragraphs[0]
        otp.text = o_ti
        otp.font.size = Pt(10.5)
        otp.font.bold = True
        otp.font.color.rgb = o_col
        otp.font.name = "Segoe UI"
        
        otp_sub = otf.add_paragraph()
        otp_sub.text = f"▶ {o_sub}"
        otp_sub.font.size = Pt(9.5)
        otp_sub.font.bold = True
        otp_sub.font.color.rgb = COLOR_WHITE
        otp_sub.font.name = "Consolas"
        
        otp_de = otf.add_paragraph()
        otp_de.text = o_de
        otp_de.font.size = Pt(8.5)
        otp_de.font.color.rgb = COLOR_TEXT_MUTED
        otp_de.font.name = "Segoe UI"

    upg_img_path = os.path.join(assets_dir, "upgrade.png")
    if os.path.exists(upg_img_path):
        slide6.shapes.add_picture(upg_img_path, Inches(8.9), Inches(2.1), Inches(3.6), Inches(3.6))
        
    add_card(slide6, 8.9, 5.85, 3.6, 1.15, COLOR_CARD_ALT, COLOR_GREEN, 1.2)
    up_tb = slide6.shapes.add_textbox(Inches(9.05), Inches(5.95), Inches(3.3), Inches(0.95))
    up_tf = up_tb.text_frame
    up_tf.word_wrap = True
    up_tf.margin_left = up_tf.margin_top = up_tf.margin_right = up_tf.margin_bottom = 0
    upp = up_tf.paragraphs[0]
    upp.text = "⚡ UPGRADE: SYSTEM SYNERGY"
    upp.font.size = Pt(10.5)
    upp.font.bold = True
    upp.font.color.rgb = COLOR_GREEN
    upp.font.name = "Consolas"
    
    upp2 = up_tf.add_paragraph()
    upp2.text = "Unified data flow ensures candidate telemetry seamlessly feeds all three predictive engines in < 1.2s total latency."
    upp2.font.size = Pt(8.5)
    upp2.font.color.rgb = COLOR_TEXT_MUTED
    upp2.font.name = "Segoe UI"

    notes6 = slide6.notes_slide.notes_text_frame
    notes6.text = (
        "[TRANSITION CUE: SLIDE 5 -> SLIDE 6]\n"
        "- Camera dives into the inner circuits of the Omnitrix.\n"
        "- 3D Upgrade attaches his glowing biomechanical tentacles directly into the 4-step pipeline.\n"
        "- Data pulses travel in real time across the glowing neon green conduits into the three predictive engines."
    )

    # =========================================================================
    # SLIDE 7: CONCLUSION & FUTURE SCOPE
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide7)
    add_header(slide7, "⚡ MODULE 06 // SUMMARY & IMPACT", "CONCLUSION & FUTURE SCOPE",
               "Transforming talent acquisition into a fast, transparent, and data-driven powerhouse.")
    
    pillars = [
        ("🚀", "SMARTER RECRUITMENT", "Semantic deep understanding matches candidates to roles based on genuine ability, cutting through keyword manipulation and pedigree bias.", COLOR_CYAN),
        ("⚡", "FASTER DECISION-MAKING", "Reduces candidate screening and shortlisting cycles from weeks to minutes, allowing talent acquisition teams to act before competitors.", COLOR_GREEN),
        ("📊", "DATA-DRIVEN EVALUATION", "Replaces subjective intuition with quantifiable readiness, predictive retention scores, and standardized competency benchmarks.", COLOR_BLUE),
        ("⚖️", "FAIR & CONSISTENT ASSESSMENTS", "Objective NLP models evaluate candidates on proven capability and growth trajectory, ensuring an equitable hiring process.", COLOR_PURPLE),
    ]
    
    for p_idx, (p_icon, p_ti, p_de, p_col) in enumerate(pillars):
        px = 0.8 + (p_idx % 2) * 4.35
        py = 2.1 + (p_idx // 2) * 2.05
        
        c = add_card(slide7, px, py, 4.15, 1.9, COLOR_CARD, p_col, 1.4)
        
        tb = slide7.shapes.add_textbox(Inches(px + 0.2), Inches(py + 0.15), Inches(3.75), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"{p_icon}  {p_ti}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = p_col
        p.font.name = "Segoe UI"
        
        p_desc = tf.add_paragraph()
        p_desc.text = p_de
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = COLOR_TEXT_MUTED
        p_desc.font.name = "Segoe UI"

    if os.path.exists(ben_img_path):
        slide7.shapes.add_picture(ben_img_path, Inches(9.6), Inches(1.9), Inches(3.0), Inches(3.9))
        
    banner = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(8.5), Inches(0.95))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(10, 24, 18)
    banner.line.color.rgb = COLOR_GREEN
    banner.line.width = Pt(1.8)
    
    b_tf = banner.text_frame
    b_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = b_tf.paragraphs[0]
    bp.text = "FROM RESUME TO SUCCESSFUL HIRING — POWERED BY AI"
    bp.font.size = Pt(15)
    bp.font.bold = True
    bp.font.color.rgb = COLOR_GREEN
    bp.font.name = "Segoe UI"
    bp.alignment = PP_ALIGN.CENTER
    
    bp_sub = b_tf.add_paragraph()
    bp_sub.text = "Empowering organizations to build high-performance, future-ready engineering teams."
    bp_sub.font.size = Pt(10)
    bp_sub.font.color.rgb = COLOR_WHITE
    bp_sub.font.name = "Segoe UI"
    bp_sub.alignment = PP_ALIGN.CENTER

    demo_btn = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.6), Inches(6.15), Inches(3.0), Inches(0.95))
    demo_btn.fill.solid()
    demo_btn.fill.fore_color.rgb = COLOR_GREEN
    demo_btn.line.fill.background()
    d_tf = demo_btn.text_frame
    d_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    dp = d_tf.paragraphs[0]
    dp.text = "▶  LIVE DEMO"
    dp.font.size = Pt(16)
    dp.font.bold = True
    dp.font.color.rgb = RGBColor(0, 0, 0)
    dp.font.name = "Segoe UI"
    dp.alignment = PP_ALIGN.CENTER
    
    dp_sub = d_tf.add_paragraph()
    dp_sub.text = "Click to launch system"
    dp_sub.font.size = Pt(9)
    dp_sub.font.bold = True
    dp_sub.font.color.rgb = RGBColor(20, 40, 20)
    dp_sub.font.name = "Segoe UI"
    dp_sub.alignment = PP_ALIGN.CENTER

    notes7 = slide7.notes_slide.notes_text_frame
    notes7.text = (
        "[FINAL CINEMATIC SCENE & CTA]\n"
        "- Ben Tennyson stands proudly, raising the glowing Omnitrix.\n"
        "- A pulse of emerald energy illuminates the entire room as the conclusion cards settle into view.\n"
        "- Click the pulsating 'LIVE DEMO' button to transition seamlessly into the running application demonstration."
    )

    output_pptx = r"c:\Users\shinc\OneDrive\Documents\AI RECRUITMENT PLATFORM\New folder\Ben10_AI_Recruitment_Presentation.pptx"
    prs.save(output_pptx)
    print(f"Presentation successfully created at: {output_pptx}")

if __name__ == "__main__":
    build_ben10_presentation()
