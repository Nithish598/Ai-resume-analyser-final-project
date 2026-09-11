"""Generate 10 Realistic, Anonymized Sample Resumes in PDF, DOCX, and TXT formats.

Covers:
1. Fresher Software Engineer (.pdf)
2. Experienced Full Stack Developer (.docx)
3. Data Scientist / ML Engineer (.pdf)
4. Missing Sections Resume (.txt)
5. Unusual Headers Resume (.docx)
6. Multi-Education Academic CV (.pdf)
7. Dense Skills DevOps Engineer (.docx)
8. No Dedicated Skills Section (.txt)
9. Irregular Spacing & Formatting (.txt)
10. Scanned Image-Only Simulation PDF (.pdf)
"""
import os
import io
import docx
from docx.shared import Pt, Inches, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from PIL import Image, ImageDraw, ImageFont


OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "sample_resumes")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def create_pdf_text_resume(filename: str, title: str, sections: list):
    """Generate clean text-based PDF using ReportLab canvas."""
    filepath = os.path.join(OUTPUT_DIR, filename)
    c = canvas.Canvas(filepath, pagesize=letter)
    width, height = letter
    
    y = height - 50
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, title)
    y -= 25
    
    for sec_title, lines in sections:
        if y < 80:
            c.showPage()
            y = height - 50
            
        if sec_title:
            c.setFont("Helvetica-Bold", 12)
            c.setFillColor(colors.HexColor("#1A365D"))
            c.drawString(50, y, sec_title)
            y -= 4
            c.setStrokeColor(colors.HexColor("#CBD5E0"))
            c.setLineWidth(0.5)
            c.line(50, y, width - 50, y)
            y -= 14
            
        c.setFont("Helvetica", 9)
        c.setFillColor(colors.black)
        for line in lines:
            if y < 60:
                c.showPage()
                y = height - 50
                c.setFont("Helvetica", 9)
            c.drawString(60 if line.startswith("•") or line.startswith("-") else 50, y, line)
            y -= 13
        y -= 8
        
    c.save()
    print(f"Generated: {filepath}")


def create_docx_resume(filename: str, title: str, contact_info: str, sections: list):
    """Generate structured DOCX file using python-docx."""
    filepath = os.path.join(OUTPUT_DIR, filename)
    doc = docx.Document()
    
    # Title / Name
    h = doc.add_heading(title, level=0)
    h.paragraph_format.space_after = Pt(2)
    
    # Contact paragraph
    p = doc.add_paragraph(contact_info)
    p.paragraph_format.space_after = Pt(12)
    
    for sec_title, content in sections:
        sec_h = doc.add_heading(sec_title, level=1)
        sec_h.paragraph_format.space_before = Pt(10)
        sec_h.paragraph_format.space_after = Pt(4)
        
        if isinstance(content, list):
            for item in content:
                bp = doc.add_paragraph(item, style='List Bullet' if not item.startswith("•") and not item.startswith("-") else 'Normal')
                bp.paragraph_format.space_after = Pt(2)
        elif isinstance(content, str):
            doc.add_paragraph(content)
            
    doc.save(filepath)
    print(f"Generated: {filepath}")


def create_txt_resume(filename: str, content: str):
    """Generate UTF-8 TXT file."""
    filepath = os.path.join(OUTPUT_DIR, filename)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Generated: {filepath}")


def create_scanned_pdf(filename: str):
    """Generate a PDF containing only a rasterized image with no text layer (scanned simulation)."""
    filepath = os.path.join(OUTPUT_DIR, filename)
    
    # Create an image using PIL
    img = Image.new("RGB", (600, 800), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([20, 20, 580, 780], outline=(200, 200, 200), width=2)
    draw.text((50, 50), "SCANNED CANDIDATE RESUME", fill=(0, 0, 0))
    draw.text((50, 80), "John Doe - Hand Signed Application", fill=(80, 80, 80))
    draw.text((50, 110), "This page was scanned on a flatbed scanner at 300 DPI.", fill=(100, 100, 100))
    
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format='JPEG')
    img_byte_arr.seek(0)
    
    # Write into PDF canvas as raw image
    c = canvas.Canvas(filepath, pagesize=letter)
    width, height = letter
    from reportlab.lib.utils import ImageReader
    c.drawImage(ImageReader(img_byte_arr), 50, 100, width=500, height=650)
    c.save()
    print(f"Generated scanned simulation: {filepath}")


def generate_all_samples():
    """Build all 10 sample resumes."""
    
    # 1. Fresher Software Engineer (PDF)
    create_pdf_text_resume(
        "01_fresher_software_engineer.pdf",
        "Alex Rivera",
        [
            ("", [
                "alex.rivera@email.com | +1 (555) 234-5678 | San Francisco, CA",
                "LinkedIn: linkedin.com/in/alexrivera-cs | GitHub: github.com/alexrivera-dev | Portfolio: https://alexrivera.dev"
            ]),
            ("PROFESSIONAL SUMMARY", [
                "Motivated Computer Science graduate with a strong foundation in data structures, algorithms, and full-stack web development.",
                "Passionate about building scalable backend services and distributed systems using Python and Java."
            ]),
            ("EDUCATION", [
                "Bachelor of Science in Computer Science",
                "University of California, Berkeley (2020 - 2024)",
                "GPA: 3.85 / 4.0 | Dean's Honor List"
            ]),
            ("TECHNICAL SKILLS", [
                "Programming Languages: Python, Java, C++, JavaScript, TypeScript, SQL",
                "Frameworks & Libraries: React, Node.js, Django, FastAPI, Flask, Express.js",
                "Databases: PostgreSQL, MongoDB, Redis, SQLite",
                "Tools & Cloud: Git, Docker, AWS EC2, AWS S3, Linux, Postman, CI/CD"
            ]),
            ("EXPERIENCE", [
                "Software Engineer Intern at Nexus Tech Labs (Jun 2023 - Aug 2023)",
                "• Designed and implemented RESTful APIs using Python FastAPI and PostgreSQL, serving 10,000+ daily requests.",
                "• Containerized microservices using Docker and automated integration tests via GitHub Actions.",
                "• Optimized database indexing queries, reducing average API response latency by 28%."
            ]),
            ("PROJECTS", [
                "DevConnect - Developer Community Platform",
                "• Built a full-stack discussion forum with real-time notifications using React, Node.js, and Redis.",
                "• GitHub: https://github.com/alexrivera-dev/devconnect",
                "",
                "Automated Code Reviewer Bot",
                "• Developed a Python CLI tool using AST parsing to detect code smells and style violations.",
                "• GitHub: https://github.com/alexrivera-dev/code-reviewer"
            ]),
            ("CERTIFICATIONS", [
                "• AWS Certified Cloud Practitioner - Amazon Web Services (2023)",
                "• Meta Full-Stack Software Engineer Certificate - Coursera (2023)"
            ]),
            ("LANGUAGES", [
                "English (Fluent), Spanish (Native)"
            ])
        ]
    )
    
    # 2. Experienced Full Stack Developer (DOCX)
    create_docx_resume(
        "02_experienced_fullstack_dev.docx",
        "Sarah Jenkins",
        "sarah.jenkins@techpros.io | +1-415-555-0199 | Austin, TX\nLinkedIn: https://linkedin.com/in/sarahjenkins-dev | GitHub: https://github.com/sjenkins-fullstack",
        [
            ("PROFESSIONAL SUMMARY", 
             "Lead Full Stack Developer with 6+ years of hands-on experience designing and scaling modern enterprise web applications. Specialized in TypeScript, React, Node.js, GraphQL, and AWS cloud architecture."),
            ("WORK EXPERIENCE", [
                "Senior Software Engineer at Horizon Cloud Systems (Jan 2021 - Present)",
                "- Architected scalable micro-frontend architecture using React, Next.js, and TypeScript for 500k active users.",
                "- Engineered high-throughput RESTful and GraphQL backend services in Node.js and NestJS backed by PostgreSQL and Redis.",
                "- Led a squad of 6 engineers, conducting code reviews and mentoring junior developers in Agile/Scrum environment.",
                "- Deployed multi-region infrastructure on AWS using Terraform, ECS, S3, and CloudFront.",
                "",
                "Full Stack Developer at Apex Digital Solutions (Jun 2018 - Dec 2020)",
                "- Developed customer onboarding portals with React, Redux, Python Django, and MySQL.",
                "- Implemented JWT-based authentication and role-based access control (RBAC).",
                "- Automated deployment pipelines using Jenkins and Docker, decreasing release turnaround by 45%."
            ]),
            ("SKILLS", [
                "Programming Languages: JavaScript, TypeScript, Python, Go, HTML5, CSS3, SQL",
                "Frameworks & Web: React, Next.js, Vue.js, Node.js, NestJS, Express.js, Django",
                "Databases: PostgreSQL, MongoDB, Redis, MySQL, DynamoDB",
                "Cloud & DevOps: AWS (EC2, S3, Lambda), Docker, Kubernetes, Terraform, CI/CD, Git"
            ]),
            ("EDUCATION", [
                "Master of Science in Software Engineering",
                "University of Texas at Austin (2016 - 2018)",
                "Bachelor of Technology in Information Technology",
                "Texas State University (2012 - 2016)"
            ]),
            ("CERTIFICATIONS", [
                "AWS Certified Solutions Architect - Associate (2022)",
                "Certified Kubernetes Application Developer (CKAD) - Linux Foundation (2021)"
            ]),
            ("LANGUAGES", [
                "English (Native), French (Intermediate)"
            ])
        ]
    )

    # 3. Data Scientist / ML Engineer (PDF)
    create_pdf_text_resume(
        "03_data_scientist_ml_engineer.pdf",
        "Dr. Marcus Vance",
        [
            ("", [
                "marcus.vance@ai-research.org | +1 (617) 555-8921 | Boston, MA",
                "LinkedIn: linkedin.com/in/marcus-vance-ml | GitHub: github.com/marcusvance-ai"
            ]),
            ("EXECUTIVE SUMMARY", [
                "Data Scientist and Machine Learning Engineer with 4 years of experience building predictive models, NLP pipelines, and computer vision systems. Proficient in PyTorch, TensorFlow, Scikit-learn, and Big Data processing."
            ]),
            ("TECHNICAL PROFICIENCIES", [
                "Machine Learning & AI: Deep Learning, NLP, Computer Vision, Transformers, LLMs, Scikit-learn",
                "Languages: Python, R, C++, SQL, Bash",
                "Frameworks & Libraries: PyTorch, TensorFlow, Keras, Pandas, NumPy, SciPy, OpenCV, Hugging Face, spaCy",
                "Databases & Big Data: Apache Spark, Hadoop, PostgreSQL, MongoDB, Snowflake, BigQuery",
                "Tools & Cloud: Docker, Git, MLflow, AWS S3, AWS SageMaker, GCP Vertex AI"
            ]),
            ("WORK EXPERIENCE", [
                "Machine Learning Engineer at BioInformatics AI Labs (Aug 2021 - Present)",
                "• Developed Transformer-based NLP models to extract clinical insights from 2M+ medical records using PyTorch and Hugging Face.",
                "• Built automated feature engineering pipelines in Apache Spark and Pandas, speeding up model training cycles by 35%.",
                "• Deployed real-time inference microservices with FastAPI and Docker on AWS SageMaker with 99.9% uptime.",
                "",
                "Data Science Intern at Cambridge Analytics (Jan 2021 - Jul 2021)",
                "• Trained computer vision classification models with OpenCV and TensorFlow, achieving 94.2% top-1 accuracy.",
                "• Conducted exploratory data analysis and visualized churn dynamics with Seaborn and Matplotlib."
            ]),
            ("EDUCATION", [
                "Ph.D. in Computer Science (Specialization in Machine Learning)",
                "Massachusetts Institute of Technology (MIT) (2017 - 2021)",
                "Bachelor of Science in Mathematics and Statistics",
                "Carnegie Mellon University (2013 - 2017)"
            ]),
            ("PROJECTS", [
                "NeuroVision: Brain Tumor Segmentation",
                "• Implemented 3D U-Net convolutional neural network in PyTorch for MRI volumetric segmentation.",
                "• GitHub: https://github.com/marcusvance-ai/neuro-vision"
            ]),
            ("ACHIEVEMENTS", [
                "• Published 3 first-author papers at NeurIPS and CVPR workshops.",
                "• 1st Place Winner - Kaggle Medical Imaging Classification Challenge (2022)."
            ])
        ]
    )

    # 4. Missing Sections Resume (TXT)
    create_txt_resume(
        "04_missing_sections_resume.txt",
        """David Chen
david.chen@mailbox.net
+1 206 555 4321
Seattle, WA

EDUCATION
Bachelor of Science in Computer Engineering
University of Washington, 2019 - 2023

WORK EXPERIENCE
Software Developer at CloudPeak Technologies (Jul 2023 - Present)
- Writing backend services in Go and Python.
- Maintaining PostgreSQL relational databases and building queries.
- Participating in daily standups and sprint planning.
"""
    )

    # 5. Unusual Section Headers Resume (DOCX)
    create_docx_resume(
        "05_unusual_headers_resume.docx",
        "Elena Rostova",
        "elena.rostova@designcode.co | +44 20 7946 0912 | London, UK\nLinkedIn: https://linkedin.com/in/elena-rostova | Portfolio: https://elenarostova.design",
        [
            ("ABOUT ME", 
             "Passionate Frontend Engineer and UI/UX Designer who loves crafting accessible, visually stunning digital experiences."),
            ("MY TOOLKIT & ARSENAL", [
                "Programming Languages: JavaScript, TypeScript, HTML5, CSS3, SASS",
                "Frameworks: React, Vue.js, Svelte, Tailwind CSS, Bootstrap",
                "Design & Prototyping: Figma, Adobe XD, Storybook",
                "Databases & Backend: Firebase, Node.js, Express.js"
            ]),
            ("CAREER STORY & JOURNEY", [
                "Frontend Specialist at PixelPerfect Agency (Mar 2022 - Present)",
                "- Built responsive, mobile-first web applications using React and Tailwind CSS.",
                "- Increased lighthouse accessibility scores from 65 to 98 across 12 client websites.",
                "",
                "Junior Web Developer at BrightSpark Media (Sep 2020 - Feb 2022)",
                "- Created custom interactive widgets using JavaScript and CSS animations."
            ]),
            ("WHERE I STUDIED", [
                "Bachelor of Arts in Digital Media and Web Development",
                "University of the Arts London (2017 - 2020)"
            ]),
            ("THINGS I HAVE BUILT", [
                "CryptoWatch Dashboard",
                "- Real-time cryptocurrency monitoring dashboard using React, WebSockets, and Chart.js.",
                "URL: https://github.com/elena-rostova/cryptowatch"
            ]),
            ("HONORS & RECOGNITION", [
                "Awwwards Site of the Day Nominee (2023)",
                "Best UX Design Award - London Hackathon 2021"
            ])
        ]
    )

    # 6. Multi-Education Academic CV (PDF)
    create_pdf_text_resume(
        "06_multi_education_academic_cv.pdf",
        "Dr. Robert Sterling",
        [
            ("", [
                "robert.sterling@oxford.ac.uk | +44 1865 270000 | Oxford, United Kingdom",
                "LinkedIn: linkedin.com/in/robertsterling-phd | GitHub: github.com/rsterling"
            ]),
            ("CAREER OBJECTIVE", [
                "Senior Research Scientist seeking research and development roles in quantum computing, cryptography, and distributed systems."
            ]),
            ("ACADEMIC QUALIFICATIONS", [
                "Doctor of Philosophy in Computer Science",
                "University of Oxford (2018 - 2022)",
                "Dissertation: Fault-Tolerant Distributed Consensus in Quantum Networks",
                "",
                "Master of Science in Advanced Computing",
                "Imperial College London (2016 - 2018)",
                "Graduated with Distinction | GPA: 4.0/4.0",
                "",
                "Bachelor of Science in Mathematics and Computer Science",
                "University of Manchester (2012 - 2016)",
                "First Class Honours (Summa Cum Laude)"
            ]),
            ("TECHNICAL EXPERTISE", [
                "Programming: C++, Rust, Python, Go, Haskell, Assembly, SQL",
                "Specializations: Distributed Systems, Cryptography, Blockchain, Quantum Computing",
                "Tools: Git, Linux, Docker, LaTeX, CMake, GDB, Valgrind"
            ]),
            ("PROFESSIONAL EXPERIENCE", [
                "Senior Research Engineer at Quantex Labs (Jan 2022 - Present)",
                "• Engineered high-performance cryptographic primitives in Rust and C++.",
                "• Published peer-reviewed research and led patent filings on post-quantum key exchange algorithms."
            ]),
            ("LANGUAGES KNOWN", [
                "English (Native), German (Fluent), Latin (Reading)"
            ])
        ]
    )

    # 7. Dense Skills DevOps Engineer (DOCX)
    create_docx_resume(
        "07_dense_skills_devops.docx",
        "Karthik Subramanian",
        "karthik.subramanian@cloudops.net | +91 98765 43210 | Bangalore, India\nLinkedIn: https://linkedin.com/in/karthik-cloud | GitHub: https://github.com/karthik-devops",
        [
            ("PROFESSIONAL SUMMARY", 
             "Lead DevOps & Cloud Infrastructure Architect with 7 years of experience orchestrating enterprise Kubernetes clusters, automated CI/CD pipelines, and cloud migration initiatives across AWS, Azure, and GCP."),
            ("CORE COMPETENCIES", [
                "Cloud Platforms: AWS, Amazon Web Services, EC2, S3, Lambda, Microsoft Azure, Google Cloud Platform (GCP)",
                "Containerization & Orchestration: Docker, Kubernetes, Helm, OpenShift, Podman",
                "Infrastructure as Code: Terraform, Ansible, CloudFormation, Pulumi",
                "CI/CD Tools: Jenkins, GitLab CI, GitHub Actions, CircleCI, ArgoCD, Spinnaker",
                "Monitoring & Logging: Prometheus, Grafana, ELK Stack, Splunk, Datadog, CloudWatch",
                "Databases: PostgreSQL, MySQL, Redis, MongoDB, Cassandra, DynamoDB",
                "Programming & Scripting: Python, Bash, Shell Scripting, Go, YAML, Groovy",
                "Networking & Security: Nginx, HAProxy, SSL/TLS, Vault, SonarQube, Cybersecurity"
            ]),
            ("WORK EXPERIENCE", [
                "Principal Cloud DevOps Architect at Infosys Technologies (May 2020 - Present)",
                "- Led migration of 120+ monolithic services to microservices on AWS EKS with zero downtime.",
                "- Built end-to-end GitOps pipelines with ArgoCD, Helm, and GitHub Actions.",
                "- Reduced cloud infrastructure monthly spend by $45,000 via AWS auto-scaling and spot instances.",
                "",
                "Senior DevOps Engineer at Wipro Digital (Jul 2017 - Apr 2020)",
                "- Maintained 50+ Dockerized microservices across hybrid cloud environments.",
                "- Built custom Prometheus metrics exporters and Grafana dashboards for 24/7 reliability monitoring."
            ]),
            ("EDUCATION", [
                "Bachelor of Technology in Electronics and Communication",
                "National Institute of Technology (NIT) Karnataka (2013 - 2017)"
            ]),
            ("CERTIFICATIONS", [
                "AWS Certified Solutions Architect - Professional (2023)",
                "Certified Kubernetes Administrator (CKA) (2022)",
                "HashiCorp Certified: Terraform Associate (2021)"
            ])
        ]
    )

    # 8. No Dedicated Skills Section (TXT)
    create_txt_resume(
        "08_no_skills_section_embedded.txt",
        """Emily Watson
emily.watson@enterprise.com | (555) 789-0123 | Chicago, IL
https://linkedin.com/in/emilywatson-pm

PROFESSIONAL SUMMARY
Experienced Technical Project Manager with 5 years driving agile software development teams to deliver enterprise cloud solutions.

EXPERIENCE
Technical Project Manager at FinTech Global Corp (2021 - Present)
- Managed cross-functional squads utilizing Agile and Scrum methodologies.
- Coordinated delivery of Python and React microservices on Amazon Web Services (AWS).
- Utilized Jira, Confluence, and Git to streamline backlog tracking and sprint reporting.
- Resolved technical bottlenecks involving PostgreSQL databases and Docker containers.

Associate Project Manager at Delta Logistics (2018 - 2021)
- Spearheaded implementation of automated inventory tracking using SQL and Power BI.
- Facilitated daily standups, sprint reviews, and sprint retrospectives for 15 engineers.

EDUCATION
Bachelor of Business Administration (BBA) in Information Systems
University of Illinois at Urbana-Champaign (2014 - 2018)

CERTIFICATIONS
- Project Management Professional (PMP) - PMI (2021)
- Certified ScrumMaster (CSM) - Scrum Alliance (2019)
"""
    )

    # 9. Irregular Spacing & Formatting (TXT)
    create_txt_resume(
        "09_irregular_spacing_formatting.txt",
        """

                 RAHUL    SHARMA
    rahul.sharma@coders.in   |   +91-9988776655   |   Delhi, India
      https://github.com/rahul-sharma-code


ABOUT   ME:
Self-taught   passionate backend developer specializing in Python,  Django,  and microservices.


E X P E R I E N C E:
Backend Developer  at  Zomato  Foods  (Jan 2022  -  Present)
   • Developed REST APIs using Django and PostgreSQL
   • Implemented caching using Redis to handle 50k rpm
   • Wrote unit tests and automated CI/CD using GitHub Actions


E D U C A T I O N:
Bachelor of Technology in Computer Science
Delhi Technological University (2018 - 2022)
CGPA: 8.4/10


S K I L L S:
Python, Django, FastAPI, PostgreSQL, Redis, Docker, Git, REST API, Linux


"""
    )

    # 10. Scanned Image-Only Simulation PDF (PDF)
    create_scanned_pdf("10_scanned_image_only_simulation.pdf")

    print("\nAll 10 sample resumes successfully created in data/sample_resumes/!")


if __name__ == "__main__":
    generate_all_samples()
